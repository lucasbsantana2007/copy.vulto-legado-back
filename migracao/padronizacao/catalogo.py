"""
Catálogo de produtos — decisões 6 a 13 e 68 do Registro de Decisões.

Entrada : catalogo_vulto_2026_REVISADO_final.xlsx (abas VERAO 25, inverno25, Verão 26 [+ bloco Inverno 26],
          Atemporais; a aba de 2024 é usada só como referência de EAN e custo)
          vulto.db (opcional: fornecedor das notas de entrada, usado só para conferência)
Saída   : catalogo_corrigido.xlsx (Catálogo, Cadastro mestre, Listas, Ocorrências, Resumo)
"""
import math
import os
import sqlite3
from collections import Counter, defaultdict

import openpyxl
import pandas as pd

from comum import (CATEGORIA_POR_PREFIXO, LETRA_COR, NOME_COR, ORDEM_TAMANHO, Registro, chave, cor_padrao,
                   digitos, dinheiro, ean_digito, ean_valido, ler_codigo_produto, montar_codigo_cd, montar_sku,
                   tamanho_padrao, vazio)

ARQUIVO = "catalogo_vulto_2026_REVISADO_final.xlsx"
ABAS = [("VERAO 25", "Verão 25", 1), ("inverno25", "Inverno 25", 2), ("Verão 26", "Verão 26", 3),
        ("Atemporais", "Atemporais", 0)]
ORDEM_COLECAO = {"Verão 25": 1, "Inverno 25": 2, "Verão 26": 3, "Inverno 26": 4, "Atemporais": 0}
CABECALHOS = {
    "sku": "codigo", "código": "codigo", "cod": "codigo",
    "produto": "nome", "nome": "nome", "descricao": "nome",
    "categoria": "categoria", "cat": "categoria", "cor": "cor", "tam": "tam", "tamanho": "tam",
    "preço": "preco", "preço (r$)": "preco", "preco": "preco", "preco venda": "preco",
    "custo": "custo", "custo_medio": "custo", "custo unit": "custo",
    "ean": "ean", "ean13": "ean", "fornecedor": "fornecedor", "obs": "obs", "lançamento": "lancamento",
}
CABECALHOS = {chave(k): v for k, v in CABECALHOS.items()}
VERMELHOS = {"FFFF0000", "FFC00000"}
AMARELO = "FFFFFF00"
MARGEM_MINIMA = 0.50                     # decisão 10
TAMANHO_CD = {"XP": "PP", "PP": "PP", "S": "P", "L": "G", "XL": "GG", "UN": "U", "UNICO": "U", "TAM UNICO": "U"}

# decisão 11 — categorias com mais de um fornecedor: fornecedor pelo tipo de peça no nome.
# Primeira palavra-chave encontrada no nome decide. CONFERIR com a aba Listas do grupo.
REGRAS_FORNECEDOR_NOME = {
    "Camisas": [("jeans", "Jeans Rio Claro"), ("oxford", "Jeans Rio Claro"), ("linho", "Malharia Serra Azul"),
                ("cuban", "Confecções Aurora"), ("xadrez", "Confecções Aurora"), ("overshirt", "Confecções Aurora")],
    "Bermudas": [("jeans", "Jeans Rio Claro"), ("sarja", "Atelier Pequi"), ("cargo", "Jeans Rio Claro"),
                 ("alfaiataria", "Jeans Rio Claro"), ("moletom", "Malharia Serra Azul")],
    "Jaquetas": [("couro eco", "Jeans Rio Claro"), ("blazer", "Malharia Serra Azul"), ("bomber", "Têxtil Horizonte"),
                 ("puffer", "Têxtil Horizonte"), ("corta-vento", "Têxtil Horizonte")],
    "Tricô": [("cardigan", "Têxtil Horizonte"), ("sueter", "Têxtil Horizonte"), ("blusa", "Têxtil Horizonte"),
              ("colete", "Têxtil Horizonte"), ("tricot", "Têxtil Horizonte")],
}


def juntar(*partes):
    return "; ".join(str(x) for x in partes if not vazio(x)) or None


def _texto(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if str(v).strip() else None


# ============================================================================ leitura com formatação
def ler_abas(caminho, reg):
    wb = openpyxl.load_workbook(caminho)
    linhas = []
    for aba, colecao, _ in ABAS:
        ws = wb[aba]
        mapa, col_atual = None, colecao
        for r in range(1, ws.max_row + 1):
            valores = [ws.cell(r, j).value for j in range(1, ws.max_column + 1)]
            primeiro = chave(valores[0]) if valores[0] is not None else ""
            if primeiro.startswith("inverno 26"):
                col_atual = "Inverno 26"                   # bloco provisório dentro da aba Verão 26
                continue
            if primeiro in ("sku", "codigo", "cod"):
                mapa = {j: CABECALHOS[chave(v)] for j, v in enumerate(valores) if v and chave(v) in CABECALHOS}
                continue
            if mapa is None or valores[0] is None:
                continue
            reg_linha = {campo: valores[j] for j, campo in mapa.items()}
            celulas = [ws.cell(r, j) for j in range(1, 10)]
            preench = {c.fill.fgColor.rgb for c in celulas if c.fill.fill_type == "solid"}
            reg_linha.update(aba=aba, colecao=col_atual, linha=r,
                             tachado=any(c.font.strike for c in celulas),
                             vermelho=bool(preench & VERMELHOS), amarelo=AMARELO in preench)
            linhas.append(reg_linha)
    referencia_2024 = {}
    ws = wb["PLANILHA ANTIGA NAO USAR"]
    for n, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        p = ler_codigo_produto(r[0])
        if p and p[3]:
            sku = montar_sku(p[0], p[1], p[2], tamanho_padrao(p[3]))
            referencia_2024[sku] = dict(preco=dinheiro(r[5]), custo=dinheiro(r[6]), ean=digitos(r[7]), nome=_texto(r[1]),
                                        codigo=_texto(r[0]), tam_original=p[3], linha=n, prefixo=p[0], numero=p[1],
                                        cor=p[2], grupo=_texto(r[2]), cor_original=_texto(r[3]))
    return linhas, referencia_2024


# ============================================================================ processamento
def processar(pasta):
    reg = Registro("Catálogo")
    brutas, ref2024 = ler_abas(os.path.join(pasta, ARQUIVO), reg)

    # ---------------------------------------------------------------- identificação e tamanhos (decisões 12/13)
    itens = []
    for b in brutas:
        cod = _texto(b.get("codigo"))
        p = ler_codigo_produto(cod)
        cor = p[2] if p else cor_padrao(b.get("cor"))
        tam_orig = _texto(b.get("tam"))
        if p is None or cor is None:
            reg.add(12, "Código não reconhecido", f"{b['aba']}!{b['linha']}", "codigo", cod, None, "Linha ignorada")
            continue
        grade = tam_orig and "/" in tam_orig
        tamanhos = tam_orig.split("/") if grade else [tam_orig or p[3]]
        if grade:
            reg.add(12, "Linha de grade expandida", f"{b['aba']}!{b['linha']}", "tam", tam_orig,
                    f"{len(tamanhos)} linhas", "Uma linha com vários tamanhos virou uma linha por tamanho")
        for t in tamanhos:
            tam = tamanho_padrao(t) or tamanho_padrao(p[3])
            if tam is None:
                reg.add(12, "Tamanho não reconhecido", f"{b['aba']}!{b['linha']}", "tam", t, None)
                continue
            t_up = chave(t).upper()
            tam_cd = t_up if t_up.isdigit() else TAMANHO_CD.get(t_up, tam)
            if not grade and str(t).strip() != tam:
                reg.add(12, "Tamanho padronizado", f"{b['aba']}!{b['linha']}", "tam", t, tam)
            codigo_cd = cod.upper() if (cod and ler_codigo_produto(cod) and cod.upper()[:3].isalpha()
                                        and "-" not in cod and " " not in cod) else None
            preco, custo = dinheiro(b.get("preco")), dinheiro(b.get("custo"))
            itens.append(dict(
                colecao=b["colecao"], aba=b["aba"], linha=b["linha"], prefixo=p[0], numero=p[1], cor=cor,
                modelo=f"{p[0]}-{p[1]}", tamanho=tam, tamanho_cd=tam_cd,
                sku=montar_sku(p[0], p[1], cor, tam), codigo_cd=codigo_cd, codigo_cd_gerado=False,
                nome=_texto(b.get("nome")), categoria=CATEGORIA_POR_PREFIXO[p[0]], nome_cor=NOME_COR[cor],
                preco=preco, custo=custo, origem_custo="Nota (catálogo)" if custo is not None else None,
                ean=None if grade else digitos(b.get("ean")), origem_ean=None,
                fornecedor=_texto(b.get("fornecedor")), origem_fornecedor="Planilha" if b.get("fornecedor") else None,
                status="Descontinuado" if (b["tachado"] or b["vermelho"]) else ("Revisar preço" if b["amarelo"] else "Ativo"),
                obs=_texto(b.get("obs")), codigo_original=cod, categoria_original=_texto(b.get("categoria")),
                cor_original=_texto(b.get("cor")), tam_original=tam_orig, preco_original=_texto(b.get("preco")),
                custo_original=_texto(b.get("custo")), ean_original=_texto(b.get("ean")),
            ))
    df = pd.DataFrame(itens)

    # ---------------------------------------------------------------- decisão 9: linhas repetidas na Atemporais
    atemp = df[df.colecao == "Atemporais"]
    remover = []
    for sku, g in atemp.groupby("sku"):
        if len(g) < 2:
            continue
        boas = g[g.custo.notna()]
        manter = (boas if len(boas) else g).sort_values("preco").index[0]
        for i in g.index.drop(manter):
            remover.append(i)
            reg.add(9, "Linha repetida apagada", f"Atemporais!{df.at[i, 'linha']}", "sku", sku, None,
                    f"Repetia a linha {df.at[manter, 'linha']} com preço R$ {df.at[i, 'preco'] - df.at[manter, 'preco']:.2f} "
                    f"maior e sem custo; mantida a original")
    df = df.drop(remover)

    # ---------------------------------------------------------------- custo e preço invertidos (apoio à decisão 10)
    for i, r in df[(df.custo.notna()) & (df.preco.notna()) & (df.custo > df.preco)].iterrows():
        if r["custo"] >= 1000 and r["custo"] / 100 < r["preco"]:     # '7636' = R$ 76,36 digitado sem vírgula
            df.at[i, "custo"] = round(r["custo"] / 100, 2)
            reg.add(10, "Custo digitado sem vírgula", f"{r['aba']}!{r['linha']}", "custo", r["custo_original"],
                    df.at[i, "custo"], "Custo maior que o preço por falta da vírgula decimal")
            continue
        df.at[i, "preco"], df.at[i, "custo"] = r["custo"], r["preco"]
        reg.add(10, "Preço e custo trocados de coluna", f"{r['aba']}!{r['linha']}", "preco/custo",
                f"preço {r['preco']} / custo {r['custo']}", f"preço {r['custo']} / custo {r['preco']}",
                "Custo maior que o preço: valores digitados em colunas trocadas")

    # ---------------------------------------------------------------- decisão 6: custos
    ordem_tam = {t: i for i, t in enumerate(ORDEM_TAMANHO)}

    def mais_proximo(cands, tam):
        cands = cands[cands.custo.notna()]
        if cands.empty:
            return None
        dist = cands.tamanho.map(lambda t: abs(ordem_tam.get(t, 0) - ordem_tam.get(tam, 0)))
        return cands.loc[dist.idxmin()]

    com_custo = df[df.custo.notna() & (df.origem_custo == "Nota (catálogo)")]
    #   Atemporais: a coluna era 'custo médio'; troca pelo último custo de nota do mesmo produto nas coleções
    for i, r in df[(df.colecao == "Atemporais") & df.custo.notna()].iterrows():
        mesmo = com_custo[(com_custo.sku == r["sku"]) & (com_custo.colecao != "Atemporais")]
        if len(mesmo):
            ult = mesmo.sort_values("colecao", key=lambda s: s.map(ORDEM_COLECAO)).iloc[-1]
            df.at[i, "obs"] = juntar(r["obs"], f"custo médio anterior R$ {r['custo']:.2f}")
            df.at[i, "custo"], df.at[i, "origem_custo"] = ult["custo"], f"Último custo de nota ({ult['colecao']})"
            reg.add(6, "Custo médio trocado pelo último custo de nota", f"Atemporais!{r['linha']}", "custo",
                    r["custo"], ult["custo"], f"Mesmo produto na {ult['colecao']}; valor antigo anotado na Obs"
                    + (" (mesmo valor)" if ult["custo"] == r["custo"] else ""))
    com_custo = df[df.custo.notna()]
    pct_categoria = (com_custo.custo / com_custo.preco).groupby(com_custo.categoria).mean()
    for i, r in df[df.custo.isna()].iterrows():
        origem = f"{r['aba']}!{r['linha']}"
        mesma = com_custo[(com_custo.colecao == r["colecao"]) & (com_custo.modelo == r["modelo"]) &
                          (com_custo.cor == r["cor"])]
        outra = com_custo[(com_custo.colecao != r["colecao"]) & (com_custo.modelo == r["modelo"])]
        outra_cor = outra[outra.cor == r["cor"]]
        ref = mais_proximo(mesma, r["tamanho"])
        if ref is not None:
            custo, como = ref["custo"], f"mesmo modelo e cor, tamanho {ref['tamanho']}"
        elif (ref := mais_proximo(outra_cor if len(outra_cor) else outra, r["tamanho"])) is not None:
            custo, como = ref["custo"], f"mesmo modelo na {ref['colecao']} ({ref['nome_cor']} {ref['tamanho']})"
        elif r["prefixo"] == "ACE" and r["preco"]:
            custo, como = round(r["preco"] * pct_categoria["Acessórios"], 2), \
                f"% médio de custo dos acessórios ({pct_categoria['Acessórios']:.1%}) × preço"
        else:
            refs = [v for k, v in ref2024.items() if k.startswith(f"VLT-{r['modelo']}-") and v["custo"] and v["preco"]]
            if refs and r["preco"]:
                pct = sum(v["custo"] / v["preco"] for v in refs) / len(refs)
                custo, como = round(r["preco"] * pct, 2), f"% de custo do modelo em 2024 ({pct:.1%}) × preço"
            else:
                reg.add(6, "Custo sem referência", origem, "custo", r["custo_original"], None, "Continua em branco")
                continue
        df.at[i, "custo"], df.at[i, "origem_custo"] = custo, "Custo de referência"
        reg.add(6, "Custo em branco estimado", origem, "custo", r["custo_original"], custo, como)

    # ---------------------------------------------------------------- decisão 7: EAN
    validos = df[df.ean.map(ean_valido)]
    ean_por_sku = validos.groupby("sku").ean.first().to_dict()
    sequencia = 1
    usados = set(validos.ean)
    for i, r in df[~df.ean.map(ean_valido)].iterrows():
        origem = f"{r['aba']}!{r['linha']}"
        if r["sku"] in ean_por_sku:
            novo, como = ean_por_sku[r["sku"]], "Copiado do mesmo SKU em outra aba"
        elif ean_valido(ref2024.get(r["sku"], {}).get("ean")):
            novo, como = ref2024[r["sku"]]["ean"], "Recuperado da planilha de 2024"
        else:
            while True:
                base = f"200{sequencia:09d}"
                sequencia += 1
                novo = base + ean_digito(base)
                if novo not in usados:
                    break
            como = "EAN interno provisório (faixa 200): trocar pelo oficial antes de vender em outros canais"
        usados.add(novo)
        df.at[i, "ean"], df.at[i, "origem_ean"] = novo, como
        reg.add(7, "EAN ausente ou incompleto", origem, "ean", r["ean_original"], novo, como)
    df["origem_ean"] = df.origem_ean.fillna("Planilha")

    # ---------------------------------------------------------------- decisão 8: 'revisar preço'
    for i, r in df[df.status == "Revisar preço"].iterrows():
        outros = df[(df.colecao == r["colecao"]) & (df.modelo == r["modelo"]) & (df.index != i)]
        if len(outros) and (outros.preco == r["preco"]).all():
            df.at[i, "status"] = "Ativo"
            reg.add(8, "Revisão de preço concluída", f"{r['aba']}!{r['linha']}", "status", "Revisar preço", "Ativo",
                    f"Preço R$ {r['preco']:.2f} igual aos outros tamanhos do modelo")
        else:
            reg.add(8, "Preço continua em revisão", f"{r['aba']}!{r['linha']}", "status", "Revisar preço",
                    "Revisar preço", "Preço diferente dos outros tamanhos")

    # ---------------------------------------------------------------- decisão 10 (revisada): margem mínima de 50%
    # A migração não muda preço de venda: a margem baixa vira alerta e o preço sugerido fica para Compras aprovar.
    df["margem"] = (df.preco - df.custo) / df.preco
    df["margem_abaixo_minimo"] = "Não"
    df["preco_sugerido_margem_50"] = None
    ativos = df[df.status == "Ativo"]
    for (col, mod), g in ativos.groupby(["colecao", "modelo"]):
        if g.custo.isna().all() or g.margem.min() >= MARGEM_MINIMA - 1e-9:
            continue
        piso = g.custo.max() / (1 - MARGEM_MINIMA)
        novo = math.ceil((piso + 0.10) / 10) * 10 - 0.10           # sempre final ,90
        novo = round(max(novo, g.preco.max()), 2)
        for i in g.index:
            antigo = df.at[i, "preco"]
            df.at[i, "margem_abaixo_minimo"] = "Sim"
            if antigo == novo:
                continue
            df.at[i, "preco_sugerido_margem_50"] = novo
            reg.add(10, "Margem abaixo de 50%: preço mantido, sugestão para Compras", f"{df.at[i, 'aba']}!{df.at[i, 'linha']}",
                    "preco", antigo, antigo, f"{mod} ({col}): sugerido R$ {novo:.2f} (+R$ {novo - antigo:.2f}), não aplicado")
    df["margem"] = ((df.preco - df.custo) / df.preco).round(4)

    # ---------------------------------------------------------------- decisão 11: fornecedor
    v25 = df[(df.colecao == "Verão 25") & df.fornecedor.notna()]
    por_modelo = v25.groupby("modelo").fornecedor.agg(lambda s: Counter(s).most_common(1)[0][0]).to_dict()
    por_categoria = v25.groupby("categoria").fornecedor.agg(lambda s: Counter(s).most_common(1)[0][0]).to_dict()
    multi = set(REGRAS_FORNECEDOR_NOME)
    for i, r in df[df.fornecedor.isna()].iterrows():
        origem = f"{r['aba']}!{r['linha']}"
        if r["modelo"] in por_modelo:
            f, como = por_modelo[r["modelo"]], "Copiado do mesmo modelo"
        elif r["categoria"] in multi:
            nome = chave(r["nome"])
            f = next((forn for chave_nome, forn in REGRAS_FORNECEDOR_NOME[r["categoria"]] if chave_nome in nome), None)
            como = "Pelo tipo de peça no nome (aba Listas)"
            if f is None:
                reg.add(11, "Fornecedor sem regra", origem, "fornecedor", None, None,
                        f"'{r['nome']}' não bate com nenhuma regra da categoria {r['categoria']}")
                continue
        elif r["categoria"] in por_categoria:
            f, como = por_categoria[r["categoria"]], "Fornecedor principal da categoria na Verão 25"
        else:
            reg.add(11, "Fornecedor sem regra", origem, "fornecedor", None, None, f"Categoria {r['categoria']}")
            continue
        df.at[i, "fornecedor"], df.at[i, "origem_fornecedor"] = f, como
        reg.add(11, "Fornecedor em branco preenchido", origem, "fornecedor", None, f, como)
    nf = fornecedor_das_notas(pasta)
    df["fornecedor_nas_notas"] = df.modelo.map(nf)
    df["fornecedor_confere_notas"] = [
        None if vazio(n) else ("Sim" if f == n else "Não") for f, n in zip(df.fornecedor, df.fornecedor_nas_notas)]

    # ---------------------------------------------------------------- decisão 13: código do CD
    em_uso = codigos_em_uso(pasta)
    df["origem_codigo_cd"] = ["Planilha" if c else None for c in df.codigo_cd]
    dono = {}
    for _, r in df[df.codigo_cd.notna()].iterrows():
        dono.setdefault(r["codigo_cd"], r["sku"])
    for i, r in df[df.codigo_cd.isna()].iterrows():
        cod = montar_codigo_cd(r["prefixo"], r["numero"], r["cor"], r["tamanho_cd"])
        if dono.get(cod, r["sku"]) != r["sku"]:
            reg.add(13, "Código do CD em conflito", f"{r['aba']}!{r['linha']}", "codigo_cd", None, cod,
                    f"Código já pertence a {dono[cod]}; não gerado")
            continue
        dono[cod] = r["sku"]
        ja_usado = cod in em_uso
        df.at[i, "codigo_cd"], df.at[i, "codigo_cd_gerado"] = cod, True
        df.at[i, "origem_codigo_cd"] = ("Gerado pelo padrão — já usado pelo CD (vendas/estoque/contagem)" if ja_usado
                                        else "Provisório — confirmar cadastro no CD")
        reg.add(13, "Código do CD gerado pelo padrão", f"{r['aba']}!{r['linha']}", "codigo_cd", None, cod,
                f"{r['colecao']} — " + ("o código já aparece nos registros do CD" if ja_usado
                                         else "confirmar cadastro no CD"))
    df["obs"] = [juntar(o, "confirmar cadastro no CD") if (g and "Provisório" in str(oc)) else o
                 for o, g, oc in zip(df.obs, df.codigo_cd_gerado, df.origem_codigo_cd)]

    # ---------------------------------------------------------------- decisão 6 (revisada): custo médio ponderado
    # Dois custos convivem: o último custo de nota (Compras, para precificar) fica em 'custo'; o custo médio
    # ponderado das notas de entrada (Financeiro, para margem e estoque valorizado) vai para a Lorenzi.
    medio = custo_medio_das_notas(pasta)
    df["custo_medio"] = df.sku.map(lambda s: medio.get(s, (None, 0))[0])
    df["origem_custo_medio"] = df.sku.map(lambda s: f"Média ponderada de {medio[s][1]} peças nas notas de entrada"
                                          if s in medio else None)
    atemp = (df.colecao == "Atemporais") & df.custo_medio.isna() & df.custo_original.notna()
    df.loc[atemp, "custo_medio"] = df.loc[atemp, "custo_original"].map(dinheiro)
    df.loc[atemp, "origem_custo_medio"] = "Coluna custo_medio da aba Atemporais"
    resto = df.custo_medio.isna()
    df.loc[resto, "custo_medio"] = df.loc[resto, "custo"]
    df.loc[resto, "origem_custo_medio"] = "Sem nota com custo: igual ao último custo"
    for _, r in df[~resto & ~atemp].iterrows():
        if r["custo"] is not None and abs(r["custo_medio"] - r["custo"]) >= 0.01:
            reg.add(6, "Custo médio ponderado (Financeiro) separado do último custo", f"{r['aba']}!{r['linha']}",
                    "custo_medio", r["custo"], round(r["custo_medio"], 2), r["origem_custo_medio"])
    df["custo_medio"] = df.custo_medio.astype(float).round(2)

    # ---------------------------------------------------------------- cadastro mestre (1 linha por SKU)
    df["ordem"] = df.colecao.map(ORDEM_COLECAO)
    mestre = (df.sort_values(["sku", "ordem"]).groupby("sku").tail(1)   # coleção mais nova; Atemporais (0) só se não houver
              .drop(columns=["ordem"]).sort_values(["modelo", "cor", "tamanho"]))
    df = df.drop(columns=["ordem"])
    # decisão 68: peça em mais de uma coleção vale o MENOR preço (REGRAS_OPERACAO: "pra não dar problema com Procon";
    # o arquivo do Diego diz que, em conflito com a planilha da Marina, vale ele)
    vigentes = df[df.status != "Descontinuado"]
    menor = vigentes.groupby("sku").preco.min().combine_first(df.groupby("sku").preco.min())
    mestre["preco_colecao_mais_nova"] = mestre.preco
    mestre["preco"] = mestre.sku.map(menor)
    for _, r in mestre[mestre.preco < mestre.preco_colecao_mais_nova - 0.001].iterrows():
        reg.add(68, "Peça em duas coleções: vale o menor preço", r["sku"], "preco", r["preco_colecao_mais_nova"],
                r["preco"], f"{r['colecao']} cobrava R$ {r['preco_colecao_mais_nova']:.2f}")
    mestre["margem"] = ((mestre.preco - mestre.custo_medio) / mestre.preco).round(4)
    colunas = ["colecao", "aba", "linha", "sku", "codigo_cd", "modelo", "nome", "categoria", "nome_cor", "tamanho",
               "preco", "custo", "custo_medio", "margem", "margem_abaixo_minimo", "preco_sugerido_margem_50",
               "origem_custo", "origem_custo_medio", "ean", "origem_ean", "fornecedor", "origem_fornecedor",
               "fornecedor_nas_notas", "fornecedor_confere_notas", "status", "obs", "origem_codigo_cd", "codigo_cd_gerado",
               "codigo_original", "categoria_original", "cor_original", "tam_original", "preco_original",
               "custo_original", "ean_original"]
    # ---------------------------------------------------------------- decisão 70: peças de 2024 que ainda vendem
    novos = pecas_2024(pasta, ref2024, set(mestre.sku), por_modelo, reg)
    if len(novos):
        mestre = pd.concat([mestre, novos], ignore_index=True).sort_values(["modelo", "cor", "tamanho"])
        df = pd.concat([df, novos.drop(columns=["preco_colecao_mais_nova"])], ignore_index=True)
    listas = listas_de_referencia(por_modelo, por_categoria)
    return dict(catalogo=df[colunas].sort_values(["colecao", "linha"]), mestre=mestre[colunas + ["preco_colecao_mais_nova"]],
                listas=listas,
                skus_2024=set(ref2024)), reg


STATUS_2024 = "Coleção 2024 (só estoque)"


def pecas_2024(pasta, ref2024, no_cadastro, por_modelo, reg):
    """Decisão 70. A aba de 2024 diz 'NÃO USAR - PREÇOS 2024', mas o próprio sistema da Vulto avisa que 'a antiga tem
    os SKUs de 2024 que ainda vendem, não tirar'. Entram no cadastro as peças de 2024 que aparecem em vendas,
    estoque, entradas ou contagem desde 2025. Preço = último preço vendido no PDV/site (o da aba 2024 está
    desatualizado); custo e EAN da aba 2024."""
    from comum import data_vulto, sku_canonico
    caminho = os.path.join(pasta, "vulto.db")
    with sqlite3.connect(caminho) as c:
        # só vendas pagas de verdade: pedidos de teste (R$ 1,00) e não pagos não são preço praticado
        vendidos = pd.read_sql("""select i.sku, i.preco_unit, i.qtd, v.criado_em from itens_venda i
                                  join vendas v on v.id = i.venda_id
                                  where i.qtd > 0 and lower(trim(coalesce(v.status, ''))) in ('pago', 'aprovado', '')
                                    and coalesce(v.pedido_externo, '') not like 'S99%'
                                    and coalesce(v.pedido_externo, '') not like 'APP-99%'
                                    and lower(coalesce(v.obs, '')) not like '%teste%'""", c)
        usados = set()
        for t in ("itens_venda", "estoque", "entradas_estoque"):
            usados |= {sku_canonico(s) for (s,) in c.execute(f"select distinct sku from {t}")}
    contagem = os.path.join(pasta, "contagem_inventario_set26.xlsx")
    if os.path.exists(contagem):
        for aba, col in (("PINHEIROS 20-09", None), ("Eldorado", "codigo"), ("batel", "sku"), ("CD - posicao WMS 28-09", "SKU")):
            try:
                tab = pd.read_excel(contagem, sheet_name=aba, header=2 if col is None else 0, dtype=str)
                usados |= {sku_canonico(s) for s in tab[col or "SKU"].dropna()}
            except (ValueError, KeyError):
                pass
    vendidos["sku"] = vendidos.sku.map(sku_canonico)
    vendidos["quando"] = vendidos.criado_em.map(data_vulto)
    vendidos["preco"] = vendidos.preco_unit.map(dinheiro)
    vendidos = vendidos.dropna(subset=["sku", "quando", "preco"])
    ref = vendidos.sku.map(lambda k: (ref2024.get(k) or {}).get("preco"))
    centavos = ref.notna() & (vendidos.preco > 20 * ref)                       # venda do app gravada em centavos
    vendidos.loc[centavos, "preco"] = vendidos.loc[centavos, "preco"] / 100
    plausivel = ref.isna() | (vendidos.preco >= 0.3 * ref)                     # descarta preço absurdo
    ultimo = vendidos[plausivel].sort_values("quando").groupby("sku").last()
    medio = custo_medio_das_notas(pasta)
    linhas = []
    for sku, r in ref2024.items():
        if sku in no_cadastro or sku not in usados:
            continue
        preco, origem_preco = r["preco"], "Aba 2024 (nunca vendida desde 2025)"
        if sku in ultimo.index:
            u = ultimo.loc[sku]
            preco = u.preco
            origem_preco = f"Último preço vendido ({u.quando:%d/%m/%Y})"
        tam = tamanho_padrao(r["tam_original"])
        t_up = str(r["tam_original"]).upper()
        modelo = f"{r['prefixo']}-{r['numero']}"
        linhas.append(dict(
            colecao="2024", aba="PLANILHA ANTIGA NAO USAR", linha=r["linha"], sku=sku,
            codigo_cd=montar_codigo_cd(r["prefixo"], r["numero"], r["cor"], t_up if t_up.isdigit() else TAMANHO_CD.get(t_up, tam)),
            modelo=modelo, nome=nome_proprio_peca(r["nome"]), categoria=CATEGORIA_POR_PREFIXO[r["prefixo"]],
            nome_cor=NOME_COR[r["cor"]], tamanho=tam, preco=round(preco, 2), custo=r["custo"],
            custo_medio=round(medio[sku][0], 2) if sku in medio else r["custo"],
            margem=round((preco - (medio[sku][0] if sku in medio else r["custo"] or 0)) / preco, 4) if preco else None,
            margem_abaixo_minimo="Não", preco_sugerido_margem_50=None, origem_custo="Aba 2024",
            origem_custo_medio=(f"Média ponderada de {medio[sku][1]} peças nas notas de entrada" if sku in medio
                                else "Aba 2024"),
            ean=r["ean"] if ean_valido(r["ean"]) else None, origem_ean="Aba 2024" if ean_valido(r["ean"]) else None,
            fornecedor=por_modelo.get(modelo), origem_fornecedor="Mesmo modelo" if modelo in por_modelo else None,
            fornecedor_nas_notas=None, fornecedor_confere_notas=None, status=STATUS_2024, obs=origem_preco,
            origem_codigo_cd="Gerado pelo padrão", codigo_cd_gerado=True, codigo_original=r["codigo"],
            categoria_original=r["grupo"], cor_original=r["cor_original"], tam_original=r["tam_original"],
            preco_original=r["preco"], custo_original=r["custo"], ean_original=r["ean"],
            preco_colecao_mais_nova=round(preco, 2)))
        reg.add(70, "Peça de 2024 que ainda vende incluída no cadastro", f"PLANILHA ANTIGA NAO USAR!{r['linha']}", "sku",
                None, sku, f"Preço R$ {preco:.2f}: {origem_preco}; preço da aba 2024 R$ {r['preco'] or 0:.2f}")
    return pd.DataFrame(linhas)


def nome_proprio_peca(nome):
    """'MOLETOM CARECA ESTONADO' -> 'Moletom Careca Estonado' (igual às abas novas)."""
    from comum import nome_proprio
    return nome_proprio(nome)


def custo_medio_das_notas(pasta):
    """Custo médio ponderado por SKU: soma(qtd × custo) / soma(qtd) das notas de entrada que têm custo."""
    caminho = os.path.join(pasta, "vulto.db")
    if not os.path.exists(caminho):
        return {}
    with sqlite3.connect(caminho) as c:
        e = pd.read_sql("select sku, qtd, custo_unit from entradas_estoque where qtd > 0 and custo_unit <> ''", c)
    from comum import sku_canonico
    e["sku"] = e.sku.map(sku_canonico)
    e["custo"] = e.custo_unit.map(dinheiro)
    e = e.dropna(subset=["sku", "custo"])
    g = e.assign(total=e.qtd * e.custo).groupby("sku").agg(total=("total", "sum"), pecas=("qtd", "sum"))
    return {s: (r.total / r.pecas, int(r.pecas)) for s, r in g.iterrows()}


def codigos_em_uso(pasta):
    """Códigos curtos que o CD já usa: vendas, estoque e entradas do vulto.db e a contagem da Eldorado."""
    usados = set()
    caminho = os.path.join(pasta, "vulto.db")
    if os.path.exists(caminho):
        with sqlite3.connect(caminho) as c:
            for t in ("itens_venda", "estoque", "entradas_estoque"):
                usados |= {s for (s,) in c.execute(f"select distinct sku from {t}") if s and s.isalnum()}
    contagem = os.path.join(pasta, "contagem_inventario_set26.xlsx")
    if os.path.exists(contagem):
        usados |= set(pd.read_excel(contagem, sheet_name="Eldorado", dtype=str)["codigo"].dropna())
    return usados


def fornecedor_das_notas(pasta):
    """Fornecedor mais frequente por modelo nas notas de entrada do vulto.db (só para conferência)."""
    caminho = os.path.join(pasta, "vulto.db")
    if not os.path.exists(caminho):
        return {}
    from comum import chave as k
    apelidos = {"serra azul": "Malharia Serra Azul", "rio claro": "Jeans Rio Claro", "pequi": "Atelier Pequi",
                "aurora": "Confecções Aurora", "horizonte": "Têxtil Horizonte", "ventura": "Acessórios Ventura"}
    cont = defaultdict(Counter)
    with sqlite3.connect(caminho) as c:
        for sku, f in c.execute("select sku, fornecedor from entradas_estoque where fornecedor not in ('', 'CD')"):
            p = ler_codigo_produto(sku)
            nome = next((v for a, v in apelidos.items() if a in k(f)), None)
            if p and nome:
                cont[f"{p[0]}-{p[1]}"][nome] += 1
    return {m: c.most_common(1)[0][0] for m, c in cont.items()}


def listas_de_referencia(por_modelo, por_categoria):
    linhas = [dict(lista="Fornecedor por modelo (Verão 25)", chave=k, valor=v) for k, v in sorted(por_modelo.items())]
    linhas += [dict(lista="Fornecedor principal da categoria (Verão 25)", chave=k, valor=v)
               for k, v in sorted(por_categoria.items())]
    linhas += [dict(lista=f"Fornecedor pelo nome — {cat}", chave=p, valor=f)
               for cat, regras in REGRAS_FORNECEDOR_NOME.items() for p, f in regras]
    linhas += [dict(lista="Tamanho", chave=k, valor=v) for k, v in
               [("36", "XP"), ("38", "P"), ("40", "M"), ("42", "G"), ("44", "GG"), ("46", "XG"), ("PP", "XP"),
                ("S", "P"), ("L", "G"), ("XL", "GG"), ("Único / UN", "U")]]
    linhas += [dict(lista="Letra da cor no código do CD", chave=NOME_COR[k], valor=v) for k, v in LETRA_COR.items()]
    return pd.DataFrame(linhas)


def destaque_catalogo(linha):
    if linha["status"] == "Descontinuado":
        return "F4CCCC"
    if linha["codigo_cd_gerado"] and "Provisório" in str(linha["origem_codigo_cd"]):
        return ("E4D7F5", "codigo_cd")                            # roxo: confirmar cadastro no CD
    if linha["origem_custo"] == "Custo de referência":
        return ("FFF2CC", "custo")
    return None


if __name__ == "__main__":
    import sys
    from comum import salvar_excel
    dados, reg = processar(sys.argv[1] if len(sys.argv) > 1 else "../vulto-legado-back/data")
    salvar_excel("catalogo_corrigido.xlsx", {"Catálogo": dados["catalogo"], "Cadastro mestre": dados["mestre"],
                                             "Listas": dados["listas"], "Ocorrências": reg.df(),
                                             "Resumo": reg.resumo()},
                 destaques={"Catálogo": destaque_catalogo, "Cadastro mestre": destaque_catalogo})
    print(reg.resumo().to_string())
