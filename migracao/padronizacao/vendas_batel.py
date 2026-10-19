"""
Vendas manuais da Batel (PDV fora do ar, 02 a 22/03/2026) — decisões 14 a 20 do Registro de Decisões.

Entrada : vendas_batel_MANUAL_marco.xlsx
Saída   : vendas_batel_corrigidas.xlsx (Vendas, Totais por semana, Ocorrências, Resumo)

Fica de fora (decisão 41, pendente): conciliar estas vendas com as relançadas no vulto.db.
"""
import os
import re
from datetime import datetime

import openpyxl
import pandas as pd

from comum import Registro, chave, cpf_mascara, cpf_valido, digitos, dinheiro, vazio

ARQUIVO = "vendas_batel_MANUAL_marco.xlsx"
ANO, MES = 2026, 3
VENDEDORES = {  # decisão 17
    "gus": "Gustavo", "gustavo": "Gustavo", "gustavo.batel": "Gustavo",
    "pati": "Patrícia W.", "patricia": "Patrícia W.", "patricia w.": "Patrícia W.",
    "fer": "Fernanda K.", "fer.batel": "Fernanda K.", "fernanda k.": "Fernanda K.",
}
PAGAMENTOS = {  # decisão 18
    "pix": "Pix", "credito": "Crédito", "credito 2x": "Crédito 2x", "debito": "Débito", "dinheiro": "Dinheiro",
}


def data_marco(v, reg, linha):
    """Decisão 14: todas as vendas são de março/2026. O Excel gravou '05/03' como 3 de maio (dia e mês trocados)."""
    if isinstance(v, datetime):
        if v.month == MES:
            return v.date()
        if v.day == MES:                                    # 2026-05-03 era 05/03
            d = datetime(ANO, MES, v.month).date()
            reg.add(14, "Dia e mês trocados", linha, "data", v.strftime("%d/%m/%Y"), d.strftime("%d/%m/%Y"))
            return d
    m = re.match(r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", str(v).strip())
    if m:
        d = datetime(ANO, int(m.group(2)), int(m.group(1))).date()
        reg.add(14, "Data gravada como texto", linha, "data", v, d.strftime("%d/%m/%Y"))
        return d
    reg.add(14, "Data ilegível", linha, "data", v, None)
    return None


def processar(pasta):
    reg = Registro("Vendas Batel")
    wb = openpyxl.load_workbook(os.path.join(pasta, ARQUIVO))
    ws = wb.worksheets[0]
    cab = next(r for r in range(1, 10) if chave(ws.cell(r, 1).value) == "data")
    vendas = []
    for r in range(cab + 1, ws.max_row + 1):
        v = [ws.cell(r, j).value for j in range(1, 11)]
        if vazio(v[0]) and vazio(v[1]):                     # linhas de total e em branco
            continue
        data = data_marco(v[0], reg, r)
        # decisão 15: valor
        valor = dinheiro(v[6])
        if isinstance(v[6], str):
            reg.add(15, "Valor gravado como texto", r, "valor", v[6], valor)
        # decisão 16: CPF com zeros à esquerda e máscara
        cpf = None
        if not vazio(v[3]):
            d = digitos(v[3])
            if len(d) < 11:
                reg.add(16, "CPF sem zeros à esquerda", r, "cpf", v[3], cpf_mascara(d.zfill(11)))
                d = d.zfill(11)
            cpf = cpf_mascara(d)
            if not cpf_valido(d):
                reg.add(16, "CPF inválido", r, "cpf", v[3], cpf, "Dígito verificador não confere")
            elif cpf != str(v[3]).strip():
                reg.add(16, "CPF formatado", r, "cpf", v[3], cpf)
        # decisão 17: vendedor(a)
        vend = VENDEDORES.get(chave(v[8]), v[8])
        if vend != v[8]:
            reg.add(17, "Apelido de vendedor unificado", r, "vendedora", v[8], vend)
        # decisão 18: pagamento
        pg = PAGAMENTOS.get(chave(v[7]), v[7])
        if pg != v[7]:
            reg.add(18, "Forma de pagamento padronizada", r, "pgto", v[7], pg)
        # decisão 19: texto riscado vira status
        obs = None if vazio(v[9]) else str(v[9]).strip()
        riscado = any(ws.cell(r, j).font.strike for j in range(1, 11))
        status = "Concluída"
        if riscado:
            status = "Estornada" if "estorn" in chave(obs) else "Cancelada"
            reg.add(19, f"Venda {status.lower()} (texto riscado)", r, "status", "riscado", status,
                    f"R$ {valor:.2f} fora dos totais")
        # decisão 20: 'lancei no vulto+' vira coluna
        lancado = "lancei no vulto+" in chave(obs)
        if not cpf:
            vulto = "Não se aplica (sem CPF)"
        elif status != "Concluída":
            vulto = "Não se aplica (cancelada)"
        else:
            vulto = "Sim" if lancado else "Pendente"
        if lancado:
            obs = None
        if vulto == "Pendente":
            reg.add(20, "Pontos Vulto+ pendentes", r, "lancado_vulto", None, "Pendente",
                    "Venda com CPF ainda não lançada no Vulto+")
        hora = None if vazio(v[1]) else str(v[1]).strip()[:5]
        vendas.append(dict(linha_planilha=r, data=data, hora=hora, cliente=None if vazio(v[2]) else v[2],
                           cpf=cpf, pecas=v[4], qtd=v[5], valor=valor, pagamento=pg, vendedora=vend,
                           status=status, conta_no_total="Sim" if status == "Concluída" else "Não",
                           lancado_vulto=vulto, obs=obs, data_original=v[0], valor_original=v[6], cpf_original=v[3],
                           vendedora_original=v[8], pagamento_original=v[7]))
    df = pd.DataFrame(vendas)
    df["semana"] = df.data.map(lambda d: f"Semana {min((d.day - 2) // 7 + 1, 3)}" if d else None)
    validas = df[df.conta_no_total == "Sim"]
    totais = (validas.groupby("semana").agg(vendas=("valor", "size"), total=("valor", "sum")).reset_index())
    totais.loc[len(totais)] = ["Março (total)", len(validas), round(validas.valor.sum(), 2)]
    return dict(vendas=df, totais=totais), reg


def destaque(linha):
    if linha["status"] != "Concluída":
        return "F4CCCC"                                        # vermelho: fora dos totais
    if linha["lancado_vulto"] == "Pendente":
        return ("FFF2CC", "lancado_vulto")                      # amarelo: lançar no Vulto+
    return None


if __name__ == "__main__":
    import sys
    from comum import salvar_excel
    dados, reg = processar(sys.argv[1] if len(sys.argv) > 1 else "../vulto-legado-back/data")
    salvar_excel("vendas_batel_corrigidas.xlsx", {"Vendas": dados["vendas"], "Totais por semana": dados["totais"],
                                                  "Ocorrências": reg.df(), "Resumo": reg.resumo()},
                 destaques={"Vendas": destaque})
    print(reg.resumo().to_string())
