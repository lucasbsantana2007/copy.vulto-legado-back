"""Funções e constantes usadas por todos os scripts de padronização."""
import math
import re
import unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from zoneinfo import ZoneInfo

import pandas as pd

FUSO = ZoneInfo("America/Sao_Paulo")
DATA_CORTE = datetime(2026, 10, 1)              # posição pedida pela diretoria: 01/10/2026, 00:00 (horário de Brasília)
INICIO_PERIODO = datetime(2025, 1, 1)           # faturamento de jan/2025 a set/2026

LOJAS = {1: "Pinheiros", 2: "Eldorado", 3: "Savassi", 4: "Batel", 5: "E-commerce", 6: "App", 7: "CD Cajamar"}
LOJA_POR_NOME = {chave_: id_ for id_, n in LOJAS.items() for chave_ in (n.lower(), f"vulto {n.lower()}")} | {
    "eldorado shop": 2, "shopping eldorado": 2, "cd": 7, "cajamar": 7}

# ----------------------------------------------------------------------------- texto


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")


def chave(s):
    """Texto para comparação: sem acento, minúsculo, sem espaços sobrando."""
    if s is None or (isinstance(s, float) and math.isnan(s)):
        return ""
    return re.sub(r"\s+", " ", sem_acento(s)).strip().lower()


def vazio(v):
    return v is None or (isinstance(v, float) and math.isnan(v)) or str(v).strip() in ("", "nan", "NaN", "None")


def digitos(v):
    if vazio(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", str(v))


def nome_proprio(s):
    """'MARINA DIAS' / 'marina dias' -> 'Marina Dias' (preposições em minúsculo)."""
    if vazio(s):
        return None
    minusc = {"da", "de", "do", "das", "dos", "e"}
    partes = re.sub(r"\s+", " ", str(s)).strip().split(" ")
    return " ".join(p.lower() if i and p.lower() in minusc else p[:1].upper() + p[1:].lower()
                    for i, p in enumerate(partes))


# ----------------------------------------------------------------------------- dinheiro


def dinheiro(v):
    """'R$ 1.234,56' / '379,80' / '129.8' / 129.9 -> float com 2 casas (ou None)."""
    if vazio(v):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(Decimal(s).quantize(Decimal("0.01"), ROUND_HALF_UP))
    except InvalidOperation:
        return None


# ----------------------------------------------------------------------------- CPF


def cpf_valido(d):
    if not re.fullmatch(r"\d{11}", d or "") or d == d[0] * 11:
        return False
    for n in (9, 10):
        s = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        if (s * 10 % 11) % 10 != int(d[n]):
            return False
    return True


def cpf_mascara(d):
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if d and len(d) == 11 else d


# ----------------------------------------------------------------------------- produto

CORES = {  # código da cor no SKU -> grafias encontradas nas planilhas
    "AZM": ["azul marinho", "az marinho", "marinho", "navy"],
    "BRA": ["branco", "bco", "white"],
    "CIN": ["cinza", "cinza mescla", "cz mescla", "mescla"],
    "JES": ["jeans escuro", "jns escuro", "dark denim"],
    "JCL": ["jeans claro", "jns claro", "light denim"],
    "BEG": ["bege", "beige", "areia"],
    "VIN": ["vinho", "bordo", "burgundy"],
    "LIL": ["lilas", "lavanda"],
    "MRR": ["marrom", "brown", "cafe"],
    "OFF": ["off white", "off-white", "offwhite"],
    "PRE": ["preto", "black", "pto"],
    "OLI": ["verde oliva", "olive", "vd oliva", "oliva"],
    "AMA": ["amarelo", "amarelo manteiga", "yellow"],
}
COR_POR_NOME = {n: cod for cod, ns in CORES.items() for n in ns}
NOME_COR = {"AZM": "Azul Marinho", "BRA": "Branco", "CIN": "Cinza Mescla", "JES": "Jeans Escuro",
            "JCL": "Jeans Claro", "BEG": "Bege", "VIN": "Vinho", "LIL": "Lilás", "MRR": "Marrom",
            "OFF": "Off White", "PRE": "Preto", "OLI": "Verde Oliva", "AMA": "Amarelo"}
# letra da cor no código do CD (ex.: CAM016GP = CAM-016, Bege, P) — padrão dos 573 códigos existentes
LETRA_COR = {"AZM": "A", "BRA": "B", "CIN": "C", "JES": "E", "BEG": "G", "VIN": "H", "JCL": "J",
             "LIL": "L", "MRR": "M", "OFF": "O", "PRE": "P", "OLI": "V", "AMA": "Y"}
COR_POR_LETRA = {v: k for k, v in LETRA_COR.items()}

CATEGORIA_POR_PREFIXO = {
    "CAL": "Calças", "CAM": "Camisetas", "MOL": "Moletons", "CMS": "Camisas", "VES": "Vestidos",
    "TOP": "Tops e Cropped", "BER": "Bermudas", "SAI": "Saias", "TRI": "Tricô", "JAQ": "Jaquetas",
    "ACE": "Acessórios",
}

# Tamanho padrão do catálogo. Numeração de calça/bermuda/saia vira letra (regra do grupo).
TAMANHO_NUMERICO = {"36": "XP", "38": "P", "40": "M", "42": "G", "44": "GG", "46": "XG"}
TAMANHO_LETRA = {"PP": "XP", "XP": "XP", "P": "P", "M": "M", "G": "G", "GG": "GG", "GGG": "GGG", "XG": "XG",
                 "S": "P", "L": "G", "XL": "GG",
                 "U": "U", "UN": "U", "UNICO": "U", "TAM UNICO": "U"}
ORDEM_TAMANHO = ["XP", "P", "M", "G", "GG", "XG", "GGG", "U"]


def tamanho_padrao(t):
    if vazio(t):
        return None
    s = sem_acento(str(t)).strip().upper()
    if isinstance(t, float) and t.is_integer():
        s = str(int(t))
    return TAMANHO_NUMERICO.get(s) or TAMANHO_LETRA.get(s)


def cor_padrao(v):
    return COR_POR_NOME.get(chave(v)) if not vazio(v) else None


def ler_codigo_produto(s):
    """Lê qualquer grafia de código e devolve (prefixo, numero, cor, tamanho_original) ou None.
    Aceita: VLT-CAL-001-OLI-38 | vlt-cal-001-oli-38 | CAL001V38 | vlt cal 001 verde oliva 38 |
            vlt ves 006 marrom (linha de grade, sem tamanho)."""
    if vazio(s):
        return None
    s = str(s).strip()
    m = re.fullmatch(r"(?i)vlt-([a-z]{3})-(\d{3})-([a-z]{3})-([a-z0-9]+)", s)
    if m:
        return m.group(1).upper(), m.group(2), m.group(3).upper(), m.group(4).upper()
    m = re.fullmatch(r"([A-Z]{3})(\d{3})([A-Z])(PP|GG|P|M|G|U|\d{2})", s.upper())
    if m and m.group(3) in COR_POR_LETRA:
        return m.group(1), m.group(2), COR_POR_LETRA[m.group(3)], m.group(4)
    k = chave(s)
    m = re.fullmatch(r"vlt ([a-z]{3}) (\d{3}) (.+?)(?: (pp|p|m|g|gg|u|\d{2}))?", k)
    if m and cor_padrao(m.group(3)):
        return m.group(1).upper(), m.group(2), cor_padrao(m.group(3)), (m.group(4) or "").upper() or None
    return None


def montar_sku(prefixo, numero, cor, tamanho):
    return f"VLT-{prefixo}-{numero}-{cor}-{tamanho}"


def montar_codigo_cd(prefixo, numero, cor, tamanho_cd):
    return f"{prefixo}{numero}{LETRA_COR[cor]}{tamanho_cd}"


# ----------------------------------------------------------------------------- EAN


def ean_digito(base12):
    s = sum(int(c) * (3 if i % 2 else 1) for i, c in enumerate(base12))
    return str((10 - s % 10) % 10)


def ean_valido(e):
    e = digitos(e)
    return len(e) == 13 and ean_digito(e[:12]) == e[12] and not e.endswith("000000")


# ----------------------------------------------------------------------------- telefone


def telefone_padrao(v):
    """'+55 DD 9XXXX-XXXX'. Sem DDD (8–9 dígitos) devolve None."""
    d = digitos(v)
    if len(d) in (12, 13) and d.startswith("55"):
        d = d[2:]
    if len(d) == 11:
        return f"+55 {d[:2]} {d[2:7]}-{d[7:]}"
    if len(d) == 10:
        return f"+55 {d[:2]} {d[2:6]}-{d[6:]}"
    return None


# ----------------------------------------------------------------------------- registro de ocorrências


class Registro:
    """Uma linha por alteração feita no dado — vira a aba 'Ocorrências' e a base do Registro de Decisões."""

    def __init__(self, dominio):
        self.dominio = dominio
        self.linhas = []

    def add(self, decisao, tipo, origem, campo, original, final, descricao=""):
        self.linhas.append(dict(dominio=self.dominio, decisao=decisao, tipo=tipo, origem=origem, campo=campo,
                                valor_original=None if vazio(original) else str(original),
                                valor_final=None if vazio(final) else str(final), descricao=descricao))

    def df(self):
        cols = ["dominio", "decisao", "tipo", "origem", "campo", "valor_original", "valor_final", "descricao"]
        return pd.DataFrame(self.linhas, columns=cols)

    def resumo(self):
        d = self.df()
        if d.empty:
            return d
        return (d.groupby(["dominio", "decisao", "tipo"], dropna=False).size()
                .reset_index(name="registros").sort_values(["decisao", "registros"], ascending=[True, False]))


# ----------------------------------------------------------------------------- Excel


def salvar_excel(caminho, abas, destaques=None):
    """abas: {nome: DataFrame}. destaques: {nome_aba: função(linha_df) -> cor de preenchimento ou None}."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    destaques = destaques or {}
    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        for nome, df in abas.items():
            df.to_excel(xw, sheet_name=nome[:31], index=False)
            ws = xw.sheets[nome[:31]]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", start_color="1F3864")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if len(df):
                ws.auto_filter.ref = ws.dimensions
            for j, col in enumerate(df.columns, 1):
                amostra = [len(str(x)) for x in df[col].head(300) if not vazio(x)]
                largura = min(max([len(str(col))] + amostra) + 2, 60)
                ws.column_dimensions[get_column_letter(j)].width = largura
                if largura >= 60:
                    for c in ws[get_column_letter(j)][1:]:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
            regra = destaques.get(nome)
            if regra:
                for i, (_, linha) in enumerate(df.iterrows(), start=2):
                    estilo = regra(linha)
                    if not estilo:
                        continue
                    cor, coluna = estilo if isinstance(estilo, tuple) else (estilo, None)
                    alvo = range(1, len(df.columns) + 1) if coluna is None else [list(df.columns).index(coluna) + 1]
                    for j in alvo:
                        ws.cell(i, j).fill = PatternFill("solid", start_color=cor)


# ----------------------------------------------------------------------------- datas do vulto.db

FORMATOS_DATA = ["%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S",
                 "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"]


def data_vulto(v):
    """Data/hora do vulto.db no horário de Brasília (datetime sem fuso).

    Decisão 44: o site, o app e a Eldorado depois de 10/02/2026 gravam em UTC (terminam em 'Z'); o resto grava
    no horário local. O relatório antigo tratava tudo como se fosse o mesmo relógio."""
    if vazio(v):
        return None
    s = str(v).strip()
    utc = s.endswith("Z")
    s = s.rstrip("Z")
    for f in FORMATOS_DATA:
        try:
            d = datetime.strptime(s, f)
        except ValueError:
            continue
        if utc:
            d = d.replace(tzinfo=timezone.utc).astimezone(FUSO).replace(tzinfo=None)
        return d
    return None


# ----------------------------------------------------------------------------- código da peça


def sku_canonico(codigo):
    """Qualquer grafia do código (VLT-CAL-001-OLI-38, CAL001V38, 'vlt cal 001 verde oliva 38') -> SKU do catálogo."""
    p = ler_codigo_produto(codigo)
    if not p or not p[3] or not tamanho_padrao(p[3]):
        return None
    return montar_sku(p[0], p[1], p[2], tamanho_padrao(p[3]))


# ----------------------------------------------------------------------------- quem é o cliente

ANONIMOS = {"", "consumidor", "nao informado", "não informado", "cliente", "-", "sem cpf", "nao tem", "não tem"}


def identificar(ref):
    """cliente_ref do vulto.db ('cpf, email ou telefone, depende de quem cadastrou') -> (tipo, valor).

    tipo: 'cpf' (11 dígitos válidos), 'email', 'telefone' ('+55 DD 9XXXX-XXXX') ou 'anonimo'.
    Um número de 11 dígitos só vira CPF se o dígito verificador conferir; senão é tratado como celular."""
    if vazio(ref) or chave(ref) in ANONIMOS:
        return "anonimo", None
    s = str(ref).strip()
    if "@" in s and not s.startswith("@"):
        return "email", s.lower()
    d = digitos(s)
    pontuado_como_cpf = bool(re.fullmatch(r"\d{3}\.?\d{3}\.?\d{3}[-.]?\d{2}", s.replace(" ", "")))
    if 9 <= len(d) <= 11 and cpf_valido(d.zfill(11)) and (pontuado_como_cpf or len(d) == 11 or "(" not in s):
        return "cpf", d.zfill(11)
    t = telefone_padrao(s)
    if t:
        return "telefone", t
    if len(d) in (8, 9):
        return "telefone", None                        # celular sem DDD: não dá para identificar
    return "anonimo", None


# ----------------------------------------------------------------------------- transferências escritas no texto

PADROES_TRANSF = [
    r"enviado (?P<loja>[a-z]+) - (?P<qtd>\d+) pc?s?",   # enviado pinheiros - 1 pçs
    r"transf p/ (?P<loja>[a-z]+) (?P<qtd>\d+) ?un",       # transf p/ savassi 1un
    r"(?P<qtd>\d+) un -> (?P<loja>[a-z]+)",               # 1 un -> batel
    r"transf (?P<loja>[a-z]+) (?P<qtd>\d+)",              # TRANSF SAVASSI 1
]


def ler_transferencia(obs):
    """Saída de transferência anotada só no texto (o campo qtd fica vazio) -> (id da loja de destino, qtd)."""
    k = chave(obs)
    for p in PADROES_TRANSF:
        m = re.fullmatch(p, k)
        if m and m.group("loja") in LOJA_POR_NOME:
            return LOJA_POR_NOME[m.group("loja")], int(m.group("qtd"))
    return None, None
