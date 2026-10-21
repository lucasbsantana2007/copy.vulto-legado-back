"""Estoque da Vulto em 01/10/2026: cada contagem física avançada até a virada.

estoque em 01/10 = contado − vendido + devolvido + recebido − enviado, com os movimentos de depois da
contagem até 30/09 (vulto.db: itens das vendas e entradas_estoque).

Entradas (só leitura):
- contagem_inventario_set26.xlsx do time (SKU, código CD, status, avaria, Batel imputada, Savassi rateada)
- contagem original da Vulto (a Savassi contada por modelo, que é o número real dela)
- novo-catalogo_vulto_2026_REVISADO_final (1).xlsx (produto e siglas de cor)
- vulto.db e vulto_padronizado.db (vendas_padronizadas: status e tipo de cada venda)

Saída: migracao/saida/estoque_vulto_01-10-2026.xlsx.

Uso: python migracao/estoque_01_10.py
"""

import os
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import caminhos
import registro

CONTAGEM_TIME = os.path.join(caminhos.ENTRADA, "contagem_inventario_set26.xlsx")
CONTAGEM_ORIGINAL = os.path.join(caminhos.DADOS, "contagem_inventario_set26.xlsx")
CATALOGO = os.path.join(caminhos.ENTRADA, "novo-catalogo_vulto_2026_REVISADO_final (1).xlsx")
PADRONIZADO = caminhos.PADRONIZADO
SAIDA = os.path.join(caminhos.SAIDA, "estoque_vulto_01-10-2026.xlsx")

VIRADA = date(2026, 10, 1)
# Local -> (id no legado, aba do time, dia da contagem, como sabemos a data). Os movimentos contam do dia seguinte
# à contagem até 30/09 (a data das entradas não tem hora: o que é do dia da contagem já está contado).
LOCAIS = {
    "Pinheiros": (1, "PINHEIROS", date(2026, 9, 20), "contada em 20/09 após o fechamento (título da aba)"),
    "Eldorado": (2, "ELDORADO", date(2026, 9, 28), "sem data na planilha: considerada 28/09, como o time"),
    "Savassi": (3, "SAVASSI", date(2026, 9, 28), "sem data na planilha: considerada 28/09, como o time"),
    "Batel": (4, "BATEL", date(2026, 9, 28), "sem data na planilha: considerada 28/09, como o time"),
    "CD Cajamar": (7, "CD", date(2026, 9, 28), "posição do WMS de 28/09 (título da aba)"),
}
LOJA_DO_ID = {v[0]: k for k, v in LOCAIS.items()}
DESTINOS = {"pinheiros": "Pinheiros", "eldorado": "Eldorado", "savassi": "Savassi", "batel": "Batel"}
TAMANHO_NUMERO = {"XP": "36", "P": "38", "M": "40", "G": "42", "GG": "44", "XG": "46"}
SFS = re.compile(r"sfs|saiu da loja pinheiros", re.I)


# ---------------------------------------------------------------- SKU


def carregar_catalogo():
    """Código CD, produto e coleção de cada SKU (aba CATÁLOGO do time) e as siglas de cor (Listas do catálogo)."""
    wb = load_workbook(CONTAGEM_TIME, data_only=True)
    catalogo = {}
    for sku, codigo, produto, colecao, status, _ in wb["CATÁLOGO"].iter_rows(min_row=2, values_only=True):
        if sku:
            catalogo[sku] = dict(codigo=codigo, produto=produto, colecao=colecao, status=status)
    cores = {}
    for linha in load_workbook(CATALOGO, data_only=True)["Listas"].iter_rows(min_row=4, values_only=True):
        if linha[0] and linha[1]:
            cores[str(linha[0]).lower()] = linha[1]
    cores.update({"olive": "OLI", "oliva": "OLI", "white": "BRA", "off white": "OFF", "jeans escuro": "JES", "jeans claro": "JCL"})
    return catalogo, cores


def canonico(sku_time: str, codigo_cd: str | None) -> str:
    """SKU do time com o tamanho de volta em número nas peças numeradas (calça e bermuda): VLT-CAL-001-OLI-P -> -38.
    O código CD do CD guarda o número (CAL001V38)."""
    m = re.search(r"(3[68]|4[0246])$", str(codigo_cd or ""))
    if m:
        return sku_time.rsplit("-", 1)[0] + "-" + m.group(1)
    return sku_time


class Normalizador:
    """Leva qualquer grafia de SKU (padrão, código do CD, a do fiscal) para o SKU canônico."""

    def __init__(self, catalogo, cores):
        self.por_codigo, self.canonicos = {}, set()
        for sku, c in catalogo.items():
            can = canonico(sku, c["codigo"])
            self.canonicos.add(can)
            if c["codigo"]:
                self.por_codigo[str(c["codigo"]).upper()] = can
            self.por_codigo[sku] = can
        self.cores = cores

    def __call__(self, bruto: str) -> str | None:
        texto = str(bruto or "").strip().upper()
        if texto in self.por_codigo:
            return self.por_codigo[texto]
        partes = texto.replace(" ", "-").split("-")
        if len(partes) < 5 or partes[0] != "VLT":
            return None
        cor = " ".join(partes[3:-1]).lower()
        sigla = partes[3] if len(partes) == 5 and len(partes[3]) == 3 else self.cores.get(cor)
        if not sigla:
            return None
        tamanho = {"PP": "XP"}.get(partes[-1], partes[-1])  # PP = XP (decisão 10 do time)
        candidato = f"VLT-{partes[1]}-{partes[2]}-{sigla}-{tamanho}"
        if candidato in self.canonicos:
            return candidato
        numerado = f"VLT-{partes[1]}-{partes[2]}-{sigla}-{TAMANHO_NUMERO.get(tamanho, tamanho)}"
        if numerado in self.canonicos:
            return numerado
        return self.por_codigo.get(candidato)


# ---------------------------------------------------------------- contagens


def avaria_estimada():
    """Decisão 54 (decisão 13 da aba DECISÕES da contagem, 19/10/2026): Eldorado, Savassi e Batel não separaram a
    avaria; aplica-se a taxa real da Pinheiros (3,3%), com o mesmo peso por SKU, limitada ao estoque da loja. A aba
    AVARIA ESTIMADA traz a avaria de cada SKU por loja; ela sai do estoque vendável (como a da Pinheiros)."""
    wb = load_workbook(CONTAGEM_TIME, data_only=True)
    if "AVARIA ESTIMADA" not in wb.sheetnames:
        return {}
    linhas = list(wb["AVARIA ESTIMADA"].iter_rows(values_only=True))
    inicio = next(n for n, linha in enumerate(linhas) if linha and linha[0] == "SKU")
    cabecalho = linhas[inicio]
    colunas = {local: cabecalho.index(f"{local.upper()} – avaria est.") for local in ("Eldorado", "Savassi", "Batel")}
    estimada = {local: {} for local in colunas}
    for linha in linhas[inicio + 1:]:
        if not linha or not str(linha[0] or "").upper().startswith("VLT-"):   # pula a linha TOTAL
            continue
        for local, c in colunas.items():
            if int(linha[c] or 0):
                estimada[local][canonico(linha[0], linha[1])] = int(linha[c])
    return estimada


def contagens(catalogo):
    wb = load_workbook(CONTAGEM_TIME, data_only=True)
    contado = {}
    for local, (_, aba, _, _) in LOCAIS.items():
        linhas = {}
        cabecalho = [str(c or "").strip().upper() for c in next(wb[aba].iter_rows(min_row=1, max_row=1, values_only=True))]
        col_obs = cabecalho.index("OBS")                     # a aba ganhou colunas de avaria depois da OBS em 19/10
        for linha in wb[aba].iter_rows(min_row=2, values_only=True):
            sku, codigo, qtd = linha[0], linha[1], linha[2]
            if not sku:
                continue
            obs = str(linha[col_obs] or "") if local != "Savassi" else "ESTOQUE SIMULADO"
            situacao = "Contado"
            if local == "Savassi":
                situacao = "Estimado (rateio por modelo)"
            elif local == "Batel" and "imput" in obs.lower():
                situacao = "Estimado (imputado)"
            linhas[canonico(sku, codigo)] = dict(qtd=int(qtd or 0), situacao=situacao, obs=obs, sku_time=sku)
        contado[local] = linhas
    # avaria estimada (decisão 54): sai do contado, que passa a ser o estoque vendável
    estimada = avaria_estimada()
    for local, por_sku in estimada.items():
        for sku, qtd in por_sku.items():
            c = contado[local].setdefault(sku, dict(qtd=0, situacao="Contado", obs="", sku_time=""))
            c["qtd"] -= qtd
            c["avaria_estimada"] = qtd
            c["situacao"] += f" · avaria estimada: {qtd} (decisão 54)"
    avaria = {canonico(l[0], l[1]): int(l[2] or 0) for l in wb["AVARIA PINHEIROS"].iter_rows(min_row=2, values_only=True) if l[0]}
    # Savassi: o que a loja contou de verdade (por modelo), na contagem original da Vulto
    original = load_workbook(CONTAGEM_ORIGINAL, data_only=True)
    savassi_modelo = {}
    for modelo, descricao, qtd in original["SAVASSI (por modelo)"].iter_rows(min_row=2, values_only=True):
        if modelo:
            savassi_modelo[str(modelo).strip().upper()] = (descricao, int(qtd or 0))
    # a avaria estimada da Savassi também sai do número oficial por modelo
    for sku, qtd in estimada.get("Savassi", {}).items():
        modelo = "-".join(sku.split("-")[1:3])
        if modelo in savassi_modelo:
            descricao, contado_modelo = savassi_modelo[modelo]
            savassi_modelo[modelo] = (descricao, contado_modelo - qtd)
    return contado, avaria, savassi_modelo


# ---------------------------------------------------------------- movimentos


def data_entrada(texto) -> date | None:
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(str(texto).strip(), formato).date()
        except ValueError:
            pass
    return None


def quantidade_no_texto(obs: str) -> int | None:
    m = re.search(r"(\d+)\s*(?:un|pç|pc|p)", obs, re.I) or re.search(r"\b(\d+)\s*$", obs)
    return int(m.group(1)) if m else None


def destino_no_texto(obs: str) -> str | None:
    baixo = obs.lower()
    return next((loja for chave, loja in DESTINOS.items() if chave in baixo), None)


def movimentos(normalizar):
    """Lista de (local, dia, sku, quantidade com sinal, tipo, origem). Sem SKU reconhecido: vai para pendências."""
    conn = sqlite3.connect(PADRONIZADO)
    movs, pendentes = [], []

    # Vendas: só as que tiraram peça da loja (concluídas, sem status, trocas e devoluções); canceladas, estornadas,
    # aguardando, duplicadas e testes não. Site e app saem do CD; ship from store (obs SFS) sai da Pinheiros.
    vendas = pd.read_sql(
        "select v.id_venda, v.data, v.loja_id, v.status, v.tipo, v.duplicada_de, v.obs, i.sku, i.qtd "
        "from vendas_padronizadas v join itens_venda i on i.venda_id = v.id_venda",
        conn,
    )
    for v in vendas.itertuples():
        dia = datetime.strptime(v.data, "%d/%m/%Y").date()
        if dia < date(2026, 9, 21) or dia >= VIRADA:
            continue
        if v.tipo == "Teste" or v.duplicada_de not in ("", None) or v.status in ("Cancelada", "Estornada", "Chargeback", "Aguardando pagamento"):
            continue
        local = LOJA_DO_ID.get(v.loja_id, "CD Cajamar")
        if v.loja_id == 5 and SFS.search(v.obs or ""):
            local = "Pinheiros"
        sku = normalizar(v.sku)
        tipo = {"Devolução": "Devolução", "Troca": "Troca (peça nova)"}.get(v.tipo, "Venda")
        if sku is None:
            pendentes.append((local, dia, v.sku, -v.qtd, tipo, f"venda {v.id_venda}"))
            continue
        movs.append((local, dia, sku, -int(v.qtd), tipo, f"venda {v.id_venda}"))

    # Entradas: compras do CD, envios do CD (anotados pelo CD como "entrada" na loja 7) e recebimentos das lojas.
    # Transferência CD -> loja vale uma vez: a loja usa o próprio registro de recebimento e, sem ele, o envio do CD.
    entradas = pd.read_sql("select * from entradas_estoque order by id", conn)
    envios, recebidos, compras_vistas = [], [], set()
    for e in entradas.itertuples():
        dia = data_entrada(e.data)
        obs = str(e.obs or "")
        sku = normalizar(e.sku)
        # Sem quantidade: o CD escreve no texto ("2 un -> savassi"); a loja não ("recebido do CD romaneio 4" é o romaneio)
        qtd = int(e.qtd) if pd.notna(e.qtd) else (quantidade_no_texto(obs) if e.loja_id == 7 else None) or 1
        if dia is None:
            continue
        if e.loja_id == 7 and e.usuario in ("marina.compras", "joao.fiscal"):
            chave = (sku, qtd, e.nf)
            if e.nf and chave in compras_vistas:  # a mesma NF lançada pela Marina e pelo fiscal
                continue
            compras_vistas.add(chave)
            if sku and date(2026, 9, 29) <= dia < VIRADA:
                movs.append(("CD Cajamar", dia, sku, qtd, "Compra (NF)", f"entradas_estoque {e.id}"))
        elif e.loja_id == 7:
            destino = destino_no_texto(obs)
            envios.append(dict(id=e.id, dia=dia, sku=sku, qtd=qtd, destino=destino, bruto=e.sku))
        else:
            origem_loja = destino_no_texto(obs)
            recebidos.append(dict(id=e.id, dia=dia, sku=sku, qtd=qtd, loja=LOJA_DO_ID.get(e.loja_id), de=origem_loja or "CD", bruto=e.sku))

    # Pareia envio do CD com recebimento da loja (mesmo SKU e loja, até 10 dias depois)
    livres = defaultdict(list)
    for r in recebidos:
        if r["de"] == "CD":
            livres[(r["sku"], r["loja"])].append(r)
    for env in envios:
        par = next((r for r in livres[(env["sku"], env["destino"])] if 0 <= (r["dia"] - env["dia"]).days <= 10 and not r.get("par")), None)
        if par:
            par["par"] = env["id"]
        env["par"] = par
    for env in envios:
        if not env["sku"] or not env["destino"]:
            if date(2026, 9, 29) <= env["dia"] < VIRADA:
                pendentes.append(("CD Cajamar", env["dia"], env["bruto"], -env["qtd"], "Envio do CD sem SKU ou destino", f"entradas_estoque {env['id']}"))
            continue
        if date(2026, 9, 29) <= env["dia"] < VIRADA:
            movs.append(("CD Cajamar", env["dia"], env["sku"], -env["qtd"], f"Envio p/ {env['destino']}", f"entradas_estoque {env['id']}"))
        if env["par"] is None:  # a loja não anotou: vale o envio do CD
            inicio = LOCAIS[env["destino"]][2]
            if inicio < env["dia"] < VIRADA:
                movs.append((env["destino"], env["dia"], env["sku"], env["qtd"], "Recebido do CD (só o CD anotou)", f"entradas_estoque {env['id']}"))
    for r in recebidos:
        if not r["loja"]:
            continue
        inicio = LOCAIS[r["loja"]][2]
        if not (inicio < r["dia"] < VIRADA):
            continue
        if not r["sku"]:
            pendentes.append((r["loja"], r["dia"], r["bruto"], r["qtd"], "Recebimento sem SKU reconhecido", f"entradas_estoque {r['id']}"))
            continue
        movs.append((r["loja"], r["dia"], r["sku"], r["qtd"], f"Recebido de {r['de']}", f"entradas_estoque {r['id']}"))
        if r["de"] == "CD" and not r.get("par") and date(2026, 9, 29) <= r["dia"]:  # o CD não anotou o envio
            movs.append(("CD Cajamar", r["dia"], r["sku"], -r["qtd"], f"Envio p/ {r['loja']} (só a loja anotou)", f"entradas_estoque {r['id']}"))
    # Movimento só conta depois da contagem do local
    movs = [m for m in movs if LOCAIS[m[0]][2] < m[1] < VIRADA]
    return movs, pendentes


# ---------------------------------------------------------------- saída


def escrever(contado, avaria, savassi_modelo, catalogo_can, movs, pendentes):
    por_local = defaultdict(lambda: defaultdict(Counter))  # local -> sku -> tipo -> qtd
    for local, _, sku, qtd, tipo, _ in movs:
        por_local[local][sku][tipo] += qtd

    wb = Workbook()
    wb.remove(wb.active)
    cabecalho = PatternFill("solid", fgColor="1F3A5F")

    def aba(nome, colunas, linhas, larguras):
        ws = wb.create_sheet(nome)
        ws.append(colunas)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cabecalho
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for linha in linhas:
            ws.append(linha)
        for i, w in enumerate(larguras, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    resumo, final = [], defaultdict(dict)
    colunas_mov = ["Venda", "Troca (peça nova)", "Devolução"]
    for local, (_, _, dia, como) in LOCAIS.items():
        linhas = []
        skus = set(contado[local]) | set(por_local[local])
        tot = Counter()
        for sku in sorted(skus):
            c = contado[local].get(sku, dict(qtd=0, situacao="Não contado", obs="", sku_time=""))
            m = por_local[local][sku]
            vendido = sum(m[t] for t in colunas_mov)
            recebido = sum(q for t, q in m.items() if t.startswith("Recebido") or t.startswith("Compra"))
            enviado = sum(q for t, q in m.items() if t.startswith("Envio"))
            em_0110 = c["qtd"] + vendido + recebido + enviado
            negativo = em_0110 < 0
            final[sku][local] = max(em_0110, 0)
            info = catalogo_can.get(sku, {})
            linhas.append([sku, info.get("codigo", ""), info.get("produto", ""), info.get("status", ""), c["qtd"], vendido, recebido,
                           enviado, max(em_0110, 0), c["situacao"] + (" · ficou negativo (vendeu mais que o contado): 0" if negativo else ""),
                           c["obs"]])
            tot.update(contado=c["qtd"], vendido=vendido, recebido=recebido, enviado=enviado, final=max(em_0110, 0),
                       zerado=-em_0110 if negativo else 0)
        aba(local, ["SKU", "Código CD", "Produto", "Status no catálogo", f"Contado ({dia:%d/%m})", "Vendido (−) / devolvido (+)",
                    "Recebido / comprado", "Enviado", "Estoque 01/10/2026", "Situação", "Obs da contagem"],
            linhas, [22, 12, 34, 20, 11, 13, 12, 10, 12, 36, 50])
        resumo.append([local, f"{dia:%d/%m/%Y}", como, tot["contado"], tot["vendido"], tot["recebido"], tot["enviado"], tot["final"],
                       tot["zerado"]])

    # Savassi por modelo: o oficial
    vendas_modelo = Counter()
    for sku, m in por_local["Savassi"].items():
        vendas_modelo["-".join(sku.split("-")[1:3])] += sum(m.values())
    linhas_sav = [[modelo, desc, qtd, vendas_modelo.get(modelo, 0), max(qtd + vendas_modelo.get(modelo, 0), 0)]
                  for modelo, (desc, qtd) in sorted(savassi_modelo.items())]
    aba("Savassi por modelo", ["Modelo", "Descrição", "Contado por modelo (sem a avaria estimada)", "Movimentos 29–30/09", "Estoque 01/10/2026 (oficial)"],
        linhas_sav, [12, 40, 14, 14, 16])

    # Por produto: SKU × local
    locais = list(LOCAIS)
    aba("Por produto", ["SKU", "Código CD", "Produto", "Coleção", "Status"] + locais
        + ["Total lojas + CD", "Avaria (Pinheiros, fora)", "Avaria estimada (Eldorado, Savassi e Batel, fora)"],
        [[sku, catalogo_can.get(sku, {}).get("codigo", ""), catalogo_can.get(sku, {}).get("produto", ""),
          catalogo_can.get(sku, {}).get("colecao", ""), catalogo_can.get(sku, {}).get("status", "")]
         + [final[sku].get(l, 0) for l in locais] + [sum(final[sku].values()), avaria.get(sku, 0),
                                                     sum(contado[l].get(sku, {}).get("avaria_estimada", 0) for l in locais)]
         for sku in sorted(final)],
        [22, 12, 34, 22, 18] + [11] * len(locais) + [12, 12, 14])

    aba("Movimentos", ["Local", "Dia", "SKU", "Quantidade", "Tipo", "Origem"],
        [[l, d, s, q, t, o] for l, d, s, q, t, o in sorted(movs, key=lambda m: (m[0], m[1]))], [12, 11, 22, 10, 34, 24])
    aba("Pendências", ["Local", "Dia", "SKU como veio", "Quantidade", "Tipo", "Origem"],
        [list(p) for p in pendentes], [12, 11, 26, 10, 34, 24])
    aba("Decisões", registro.COLUNAS, registro.decisoes(27, 37) + registro.decisoes(54, 54), registro.LARGURAS)
    for linha in wb["Decisões"].iter_rows(min_row=2):
        for c in linha:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Resumo", 0)
    ws.append(["Estoque da Vulto em 01/10/2026"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Local", "Contagem", "Data da contagem", "Contado", "Vendido (−) / devolvido (+)", "Recebido / comprado", "Enviado",
               "Estoque 01/10/2026", "Peças zeradas (vendeu mais que o contado)"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cabecalho
        c.alignment = Alignment(wrap_text=True)
    for linha in resumo:
        ws.append(linha)
    total = [sum(l[i] for l in resumo) for i in range(3, 9)]
    ws.append(["Total", "", "", *total])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    oficial_sav = sum(l[4] for l in linhas_sav)
    ws.append(["Savassi oficial (por modelo, o que foi contado, sem a avaria estimada)", "", "", sum(q for _, q in savassi_modelo.values()), "", "", "", oficial_sav])
    ws.append(["Avaria da Pinheiros (fora do estoque vendável)", "", "", sum(avaria.values())])
    ws.append(["Avaria estimada de Eldorado, Savassi e Batel (fora do estoque vendável, decisão 54)", "", "",
               sum(c.get("avaria_estimada", 0) for linhas in contado.values() for c in linhas.values())])
    ws.append(["Movimentos sem SKU reconhecido (ver Pendências)", "", "", len(pendentes)])
    for i, w in enumerate([44, 12, 50, 10, 14, 14, 10, 14, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(SAIDA)
    return resumo, oficial_sav, total


def main():
    catalogo, cores = carregar_catalogo()
    catalogo_can = {canonico(sku, c["codigo"]): c for sku, c in catalogo.items()}
    normalizar = Normalizador(catalogo, cores)
    contado, avaria, savassi_modelo = contagens(catalogo)
    movs, pendentes = movimentos(normalizar)
    resumo, oficial_sav, total = escrever(contado, avaria, savassi_modelo, catalogo_can, movs, pendentes)
    for linha in resumo:
        print(f"  {linha[0]:11} contado {linha[3]:5}  vendido {linha[4]:5}  recebido {linha[5]:4}  enviado {linha[6]:5}  "
              f"01/10 {linha[7]:5}  zerado {linha[8]}")
    print("total:", total, "| savassi oficial por modelo:", oficial_sav, "| pendências:", len(pendentes))
    print("movimentos por tipo:", dict(Counter(m[4].split(" (")[0] for m in movs)))
    print("gravado:", SAIDA)


if __name__ == "__main__":
    main()
